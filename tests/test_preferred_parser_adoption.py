"""Synthetic files exercise installed preferred parsers without books or models."""
import sys
from types import SimpleNamespace

import pytest

from book_to_skill.parsers import docx as parser
from book_to_skill.parsers.epub import extract_with_ebooklib
from book_to_skill.sanitize import is_invisible_codepoint, sanitize_extracted_text


def test_precise_carriers_and_visible_neighbours():
    from tools.scan_generated_skill import _is_invisible
    carriers = [0x180B, 0x180C, 0x180D, 0x180F, 0x17B4, 0x17B5,
                0x1BCA0, 0x1BCA1, 0x1BCA2, 0x1BCA3]
    for cp in carriers:
        assert is_invisible_codepoint(cp) and _is_invisible(cp)
        assert sanitize_extracted_text('a' + chr(cp) + 'b') == ('ab', 1)
    for cp in [0x180A, 0x17B6, 0x1BC9F, 0x0600, 0x06DD, 0x08E2,
               0x110BD, 0x110CD, 0x13430, 0x13431]:
        assert not is_invisible_codepoint(cp)
        assert sanitize_extracted_text(chr(cp)) == (chr(cp), 0)
    for text in ['ᠮᠣᠩᠭᠣᠯ ᠪᠢᠴᠢᠭ', 'ជំពូក ១']:
        assert sanitize_extracted_text(text) == (text, 0)


def _docx(tmp_path, with_sdt=False):
    docx = pytest.importorskip('docx')
    document = docx.Document()
    document.add_paragraph('Before')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Cell A'
    table.cell(0, 1).text = 'Cell B'
    if with_sdt:
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        block = parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtContent>'
                          '<w:p><w:r><w:t>Inside SDT</w:t></w:r></w:p>'
                          '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>SDT table</w:t>'
                          '</w:r></w:p></w:tc></w:tr></w:tbl>'
                          '</w:sdtContent></w:sdt>')
        document._element.body.insert(-1, block)
    document.add_paragraph('After')
    path = tmp_path / 'synthetic.docx'
    document.save(path)
    return str(path)


def test_real_docx_interleaved_table(tmp_path):
    path = _docx(tmp_path)
    assert parser.extract_docx_with_python_docx(path) == 'Before\nCell A\tCell B\nAfter'


def test_real_docx_sdt_coverage_and_order(tmp_path):
    text, method = parser.extract_docx(_docx(tmp_path, True))
    assert method == 'zipfile-docx'
    assert text == 'Before\nCell A\tCell B\nInside SDT\nSDT table\nAfter'


@pytest.mark.parametrize(('outside', 'contents'), [
    ('Repeated', ['Repeated']),
    ('Repeated', ['Repeated', 'Repeated']),
    ('Prefix Repeated Suffix', ['Repeated']),
])
def test_docx_sdt_duplicate_or_substring_preserves_every_block(tmp_path, outside, contents):
    docx = pytest.importorskip('docx')
    from docx.oxml import OxmlElement

    document = docx.Document()
    document.add_paragraph(outside)
    for content in contents:
        sdt = OxmlElement('w:sdt')
        wrapper = OxmlElement('w:sdtContent')
        paragraph = OxmlElement('w:p')
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = content
        run.append(text)
        paragraph.append(run)
        wrapper.append(paragraph)
        sdt.append(wrapper)
        document._element.body.insert(-1, sdt)
    document.add_paragraph('After')
    path = tmp_path / 'duplicate-sdt.docx'
    document.save(path)

    assert parser.extract_docx_with_python_docx(str(path)) is None
    assert parser.extract_docx(str(path)) == (
        '\n'.join([outside, *contents, 'After']), 'zipfile-docx',
    )


def test_old_docx_iterator_uses_safe_fallback(tmp_path, monkeypatch):
    path = _docx(tmp_path)
    import docx
    monkeypatch.setattr(docx, 'Document', lambda _: SimpleNamespace(paragraphs=[], tables=[]))
    assert parser.extract_docx(path) == ('Before\nCell A\tCell B\nAfter', 'zipfile-docx')


def test_docx_inline_sdt_text_uses_safe_fallback(tmp_path):
    docx = pytest.importorskip('docx')
    from docx.oxml import OxmlElement

    document = docx.Document()
    paragraph = document.add_paragraph('Before ')
    sdt = OxmlElement('w:sdt')
    wrapper = OxmlElement('w:sdtContent')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = 'Inside'
    run.append(text)
    wrapper.append(run)
    sdt.append(wrapper)
    paragraph._p.append(sdt)
    document.add_paragraph('After')
    path = tmp_path / 'inline-sdt.docx'
    document.save(path)

    assert parser.extract_docx_with_python_docx(str(path)) is None
    assert parser.extract_docx(str(path)) == ('Before Inside\nAfter', 'zipfile-docx')


def test_epub_actual_spine_then_remaining_once(tmp_path):
    epub = pytest.importorskip('ebooklib.epub')
    pytest.importorskip('bs4')
    book = epub.EpubBook()
    book.set_identifier('synthetic-order')
    book.set_title('Synthetic order')
    book.set_language('en')
    items = []
    for n in [3, 1, 2]:
        item = epub.EpubHtml(uid=f'c{n}', title=f'Heading {n}', file_name=f'c{n}.xhtml')
        item.content = f'<h2>Chapter <span>{n}</span>: Topic</h2><p>hyper<span>text</span></p>'
        book.add_item(item)
        items.append(item)
    book.spine = [items[1], items[2]]
    book.toc = tuple(items)
    book.add_item(epub.EpubNav())
    path = tmp_path / 'synthetic.epub'
    epub.write_epub(str(path), book)
    text = extract_with_ebooklib(str(path))
    assert text.index('Chapter 1') < text.index('Chapter 2') < text.index('Chapter 3')
    assert text.count('Chapter 3') == 1 and 'hypertext' in text


def test_spine_invalid_refs_duplicate_and_no_spine(monkeypatch):
    pytest.importorskip('bs4')
    items = [SimpleNamespace(get_type=lambda: 9, get_content=lambda: b'<p>A</p>'),
             SimpleNamespace(get_type=lambda: 9, get_content=lambda: b'<p>B</p>')]
    book = SimpleNamespace(spine=[('b', 'yes'), 'b', [], 2, 'missing', 'image'],
                           get_items_of_type=lambda _: [items[0], items[1], items[0]],
                           get_item_with_id=lambda key: {'b': items[1], 'image': SimpleNamespace(get_type=lambda: 1)}.get(key))
    fake = SimpleNamespace(ITEM_DOCUMENT=9, epub=SimpleNamespace(read_epub=lambda _: book))
    monkeypatch.setitem(sys.modules, 'ebooklib', fake)
    monkeypatch.setitem(sys.modules, 'ebooklib.epub', fake.epub)
    assert extract_with_ebooklib('unused') == 'B\n\nA'
    book.spine = []
    assert extract_with_ebooklib('unused') == 'A\n\nB'
